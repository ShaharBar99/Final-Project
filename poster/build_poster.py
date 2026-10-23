"""Builds the A0 poster (poster/poster.pptx) from the result tables.

Run from the project root:
    python poster/build_poster.py
The charts are native PowerPoint charts filled from results/tables/, so the
poster can be edited in PowerPoint and always shows the numbers of the last run.
"""
from pathlib import Path

import pandas as pd
import qrcode
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Mm, Pt

ROOT = Path(__file__).resolve().parent.parent
TABLES = ROOT / "results" / "tables"
ASSETS = ROOT / "poster" / "assets"
REPO_URL = "https://github.com/ShaharBar99/Final-Project"

# ---------------------------------------------------------------- style

HEAD_FONT = "Cambria"
BODY_FONT = "Calibri"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


INK = "13293D"        # titles, dark panels
TEXT = "1F2A33"       # body text
MUTED = "5B6770"      # captions
PANEL = "EEF2F5"      # light panels
LINE = "C9D2D9"       # gridlines, thin borders
WHITE = "FFFFFF"

# one colour per method, used in every chart and in the legend
METHOD_STYLE = {
    "rf_without_pagerank": ("Random Forest (no PageRank input)", "D98E04", 7.0),
    "rf_with_pagerank": ("Random Forest (with PageRank input)", "F2C57C", 4.5),
    "pagerank": ("PageRank", "B23A48", 4.5),
    "closeness": ("Closeness", "5B8E7D", 3.5),
    "degree": ("Degree", "7A8B99", 3.5),
}

# order in which lines are drawn: the main model last, so it is never hidden
DRAW_ORDER = ["degree", "closeness", "rf_with_pagerank", "pagerank", "rf_without_pagerank"]

W, H = Mm(841), Mm(1189)
MARGIN = Inches(1.2)
GAP = Inches(0.8)
CONTENT_W = W - 2 * MARGIN


# ---------------------------------------------------------------- helpers

def emu(*values):
    """PowerPoint only accepts whole EMU numbers; dividing lengths in Python gives floats."""
    return [int(round(v)) for v in values]


def box(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, name=None):
    s = slide.shapes.add_shape(shape, *emu(x, y, w, h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.fill.background()
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.06
    if name:
        s.name = name
    return s


def text(slide, x, y, w, h, paragraphs, anchor=MSO_ANCHOR.TOP, name=None):
    """paragraphs: list of (runs, options); runs: list of (text, font options)."""
    tb = slide.shapes.add_textbox(*emu(x, y, w, h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, (runs, opts) in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = opts.get("align", PP_ALIGN.LEFT)
        p.space_after = Pt(opts.get("after", 0))
        p.line_spacing = opts.get("line", 1.05)
        if opts.get("bullet"):
            _bullet(p, opts.get("bullet_color", INK))
        for t, f in runs:
            r = p.add_run()
            r.text = t
            font = r.font
            font.name = f.get("font", BODY_FONT)
            font.size = Pt(f.get("size", 28))
            font.bold = f.get("bold", False)
            font.italic = f.get("italic", False)
            font.color.rgb = rgb(f.get("color", TEXT))
    return tb


def _bullet(paragraph, color):
    """Real bullet (not a typed character), with a hanging indent."""
    from pptx.oxml.ns import qn
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.45)))
    pPr.set("indent", str(-Inches(0.45)))
    for tag in ("a:buClr", "a:buFont", "a:buChar"):
        el = pPr.makeelement(qn(tag), {})
        if tag == "a:buClr":
            clr = el.makeelement(qn("a:srgbClr"), {"val": color})
            el.append(clr)
        elif tag == "a:buFont":
            el.set("typeface", "Arial")
        else:
            el.set("char", "•")
        pPr.append(el)


def keep_together(t):
    """Non-breaking space before %, non-breaking hyphen in k-core and top-10."""
    return t.replace(" %", " %").replace("k-core", "k‑core").replace("top-10", "top‑10")


def run(t, **f):
    return (keep_together(t), f)


def section_title(slide, x, y, number, title, w):
    """Numbered circle + heading - the motif of the poster."""
    d = Inches(1.05)
    c = box(slide, x, y, d, d, INK, MSO_SHAPE.OVAL, name=f"section {number} marker")
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(number)
    r.font.name, r.font.size, r.font.bold = HEAD_FONT, Pt(40), True
    r.font.color.rgb = rgb(WHITE)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    text(slide, x + d + Inches(0.35), y, w - d - Inches(0.35), d,
         [([run(title, font=HEAD_FONT, size=50, bold=True, color=INK)], {})],
         anchor=MSO_ANCHOR.MIDDLE, name=f"section {number} title")
    return y + d + Inches(0.4)


def style_axis_text(axis, size=22):
    axis.tick_labels.font.size = Pt(size)
    axis.tick_labels.font.name = BODY_FONT
    axis.tick_labels.font.color.rgb = rgb(MUTED)
    axis.format.line.color.rgb = rgb(LINE)


# ---------------------------------------------------------------- data

ranks = pd.read_csv(TABLES / "ranking_stability.csv")
cores = pd.read_csv(TABLES / "core_summary.csv")


def spearman(kind, method, level):
    row = ranks[(ranks.kind == kind) & (ranks.method == method) & (ranks.level == level)]
    return float(row["spearman_mean"].iloc[0])


def stay_top(kind, method, level=0.20):
    row = cores[(cores.kind == kind) & (cores.method == method) & (cores.level == level)]
    return float(row["important_stay_in_top"].iloc[0])


# ---------------------------------------------------------------- build

def build():
    ASSETS.mkdir(parents=True, exist_ok=True)
    qr_path = ASSETS / "qr_github.png"
    qrcode.make(REPO_URL, border=1, box_size=20).save(qr_path)

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)

    # ---------- header
    header_h = Inches(5.6)
    box(slide, 0, 0, W, header_h, INK, name="header")
    text(slide, MARGIN, Inches(0.75), Inches(25.4), Inches(3.0), [
        ([run("Stability and Structural Validity of Learned Node Importance in Citation Networks",
              font=HEAD_FONT, size=88, bold=True, color=WHITE)], {"line": 0.95}),
    ], name="title")
    text(slide, MARGIN, Inches(3.8), Inches(25.4), Inches(1.5), [
        ([run("Niran Czuckermann  ·  Shahar Bar", size=34, bold=True, color=WHITE),
          run("      Advisor: Dr. Dvora Toledano-Kitai", size=34, color="CADCE8")], {"after": 6}),
        ([run("Software Engineering Department, Braude College of Engineering  ·  Final Project 26-2-R-2",
              size=28, color="CADCE8")], {}),
    ], name="authors")
    qr_size = Inches(3.7)
    qr_x = W - MARGIN - qr_size
    box(slide, qr_x - Inches(0.2), Inches(0.5), qr_size + Inches(0.4), qr_size + Inches(0.4), WHITE,
        MSO_SHAPE.ROUNDED_RECTANGLE, name="qr frame")
    slide.shapes.add_picture(str(qr_path), *emu(qr_x, Inches(0.7), qr_size, qr_size))
    text(slide, qr_x - Inches(0.6), Inches(4.75), qr_size + Inches(1.2), Inches(0.6), [
        ([run("Code, data & live demo", size=24, bold=True, color=WHITE)], {"align": PP_ALIGN.CENTER}),
    ], name="qr caption")

    # ---------- key finding band
    y = header_h + Inches(0.5)
    band_h = Inches(4.2)
    box(slide, MARGIN, y, CONTENT_W, band_h, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE, name="key finding panel")
    pad = Inches(0.6)
    statement_w = Inches(11.6)
    text(slide, MARGIN + pad, y + pad, statement_w, band_h - 2 * pad, [
        ([run("KEY FINDING", size=26, bold=True, color="D98E04")], {"after": 10}),
        ([run("Combining centralities with a Random Forest makes the ranking more robust than "
              "PageRank — the measure it learns — against ", size=34, color=INK),
          run("added", size=34, bold=True, color=INK),
          run(" citations, but not against ", size=34, color=INK),
          run("missing", size=34, bold=True, color=INK),
          run(" ones.", size=34, color=INK)], {"line": 1.0}),
    ], anchor=MSO_ANCHOR.MIDDLE, name="key finding")

    stats = [
        (f"{stay_top('add_dag', 'rf_without_pagerank'):.0%} vs {stay_top('add_dag', 'pagerank'):.0%}",
         "of the top-10 % papers stay on top: forest vs PageRank, after adding 20 % citations"),
        (f"{spearman('remove', 'rf_without_pagerank', 0.2):.2f} vs {spearman('remove', 'pagerank', 0.2):.2f}",
         "Spearman with the original ranking after removing 20 % of citations: forest vs PageRank"),
        ("270", "damaged networks: 3 kinds of damage × 3 levels × 30 runs"),
    ]
    stat_x = MARGIN + pad + statement_w + Inches(0.6)
    stat_w = (MARGIN + CONTENT_W - pad - stat_x - 2 * Inches(0.4)) / 3
    for i, (big, small) in enumerate(stats):
        sx = stat_x + i * (stat_w + Inches(0.4))
        box(slide, sx, y + Inches(0.45), stat_w, band_h - Inches(0.9), WHITE, MSO_SHAPE.ROUNDED_RECTANGLE,
            name=f"stat card {i + 1}")
        text(slide, sx + Inches(0.3), y + Inches(0.6), stat_w - Inches(0.6), Inches(1.2), [
            ([run(big, font=HEAD_FONT, size=52, bold=True, color="D98E04" if i < 2 else INK)],
             {"align": PP_ALIGN.CENTER}),
        ], anchor=MSO_ANCHOR.MIDDLE, name=f"stat {i + 1} number")
        text(slide, sx + Inches(0.3), y + Inches(1.85), stat_w - Inches(0.6), Inches(1.7), [
            ([run(small, size=22, color=TEXT)], {"align": PP_ALIGN.CENTER, "line": 1.0}),
        ], name=f"stat {i + 1} label")

    # ---------- row 1: question / data / damage
    y = y + band_h + Inches(0.6)
    col3 = (CONTENT_W - 2 * GAP) / 3
    xs = [MARGIN + i * (col3 + GAP) for i in range(3)]
    row_top = y

    yy = section_title(slide, xs[0], row_top, 1, "The question", col3)
    text(slide, xs[0], yy, col3, Inches(4.2), [
        ([run("Centrality measures rank papers by importance, but citation data is incomplete and noisy. "
              "A small change in the network can reorder the ranking.", size=27)], {"after": 12}),
        ([run("Is a learned combination of measures more stable than any single one?", size=27, bold=True,
              color=INK)], {"after": 12}),
        ([run("Are the most important papers also the most stable?", size=27)], {"bullet": True, "after": 4}),
        ([run("Which papers look important but are fragile?", size=27)], {"bullet": True}),
    ], name="question text")

    yy = section_title(slide, xs[1], row_top, 2, "The data", col3)
    half = (col3 - Inches(0.4)) / 2
    for i, (num, label) in enumerate([("19,717", "papers on diabetes"), ("44,324", "citations")]):
        bx = xs[1] + i * (half + Inches(0.4))
        box(slide, bx, yy, half, Inches(1.9), PANEL, MSO_SHAPE.ROUNDED_RECTANGLE, name=f"data stat {i + 1}")
        text(slide, bx, yy + Inches(0.2), half, Inches(1.0), [
            ([run(num, font=HEAD_FONT, size=50, bold=True, color=INK)], {"align": PP_ALIGN.CENTER})],
            name=f"data number {i + 1}")
        text(slide, bx, yy + Inches(1.2), half, Inches(0.6), [
            ([run(label, size=24, color=MUTED)], {"align": PP_ALIGN.CENTER})], name=f"data label {i + 1}")
    text(slide, xs[1], yy + Inches(2.3), col3, Inches(2.4), [
        ([run("PubMed-Diabetes citation network. Cleaning removed 3 self-citations and 11 mutual "
              "citations: the result has no citation loops and is one connected network.", size=25)], {}),
    ], name="data text")

    yy = section_title(slide, xs[2], row_top, 3, "Damaging the network", col3)
    damage = [
        ("−", "Remove", "5, 10 or 20 % of citations — missing data"),
        ("+", "Add, newer → older", "fake citations that are still possible in time"),
        ("±", "Add at random", "the Phase A design; creates impossible citation loops"),
    ]
    for i, (sym, name, desc) in enumerate(damage):
        dy = yy + i * Inches(1.35)
        d = Inches(1.0)
        c = box(slide, xs[2], dy, d, d, PANEL, MSO_SHAPE.OVAL, name=f"damage icon {i + 1}")
        c.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = c.text_frame.paragraphs[0].add_run()
        r.text = sym
        r.font.size, r.font.bold, r.font.name = Pt(40), True, BODY_FONT
        r.font.color.rgb = rgb(INK)
        c.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        text(slide, xs[2] + d + Inches(0.3), dy - Inches(0.05), col3 - d - Inches(0.3), Inches(1.3), [
            ([run(name, size=27, bold=True, color=INK)], {}),
            ([run(desc, size=23, color=TEXT)], {"line": 1.0}),
        ], name=f"damage text {i + 1}")

    # ---------- row 2: method
    y = row_top + Inches(6.1)
    yy = section_title(slide, MARGIN, y, 4, "How we measure stability", CONTENT_W)
    steps = [
        ("Centralities", "degree, betweenness, closeness, PageRank, k-core"),
        ("Frozen scaler", "Min-Max, fitted once on the original network"),
        ("Random Forest", "learns PageRank from the other four measures"),
        ("Damage", "remove or add citations; 30 runs per setting"),
        ("Re-score", "same scaler and same forest — no retraining"),
        ("Compare", "Spearman, Kendall, top-10 % kept, per-paper CV"),
    ]
    arrow_w = Inches(0.55)
    step_gap = Inches(0.25)
    step_w = (CONTENT_W - 5 * (arrow_w + 2 * step_gap)) / 6
    step_h = Inches(2.7)
    for i, (title, desc) in enumerate(steps):
        sx = MARGIN + i * (step_w + arrow_w + 2 * step_gap)
        fill = INK if i == 2 else PANEL
        fg = WHITE if i == 2 else INK
        box(slide, sx, yy, step_w, step_h, fill, MSO_SHAPE.ROUNDED_RECTANGLE, name=f"step {i + 1}")
        text(slide, sx + Inches(0.3), yy + Inches(0.3), step_w - Inches(0.6), step_h - Inches(0.6), [
            ([run(f"{i + 1}  ", size=26, bold=True, color="D98E04"), run(title, size=28, bold=True, color=fg)],
             {"after": 8}),
            ([run(desc, size=23, color=fg if i == 2 else TEXT)], {"line": 1.0}),
        ], name=f"step {i + 1} text")
        if i < 5:
            a = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, *emu(sx + step_w + step_gap, yy + step_h / 2 - Inches(0.3),
                                                                  arrow_w, Inches(0.6)))
            a.fill.solid()
            a.fill.fore_color.rgb = rgb(LINE)
            a.line.fill.background()
            a.shadow.inherit = False
            a.name = f"arrow {i + 1}"
    text(slide, MARGIN, yy + step_h + Inches(0.3), CONTENT_W, Inches(0.8), [
        ([run("Two forests are compared with the five single measures. The forest that also gets PageRank as an "
              "input (the original design) just copies it — 99.8 % of its feature importance — so the "
              "forest without it is the main model.", size=24, italic=True, color=MUTED)], {}),
    ], name="method note")

    # ---------- row 3: Spearman charts
    y = yy + step_h + Inches(1.3)
    yy = section_title(slide, MARGIN, y, 5, "How much does the ranking change?", CONTENT_W)
    legend_h = Inches(0.6)
    _legend(slide, MARGIN, yy, CONTENT_W, legend_h)
    chart_y = yy + legend_h + Inches(0.15)
    chart_h = Inches(5.6)
    chart_w = (CONTENT_W - 2 * Inches(0.4)) / 3
    kinds = [("remove", "Citations removed"), ("add_dag", "Citations added, newer → older"),
             ("add_random", "Citations added at random")]
    for i, (kind, title) in enumerate(kinds):
        _spearman_chart(slide, MARGIN + i * (chart_w + Inches(0.4)), chart_y, chart_w, chart_h, kind, title)
    text(slide, MARGIN, chart_y + chart_h + Inches(0.15), CONTENT_W, Inches(0.9), [
        ([run("Mean Spearman correlation between the original ranking and the damaged one (30 runs per point; "
              "the runs agree closely). The forest beats PageRank when citations are added, not when they are "
              "removed. Left panel: PageRank, closeness and the forest with PageRank input lie on top of each "
              "other. k-core (not shown) behaves like degree.", size=24, color=MUTED)], {}),
    ], name="spearman caption")

    # ---------- row 4: top retention + hypotheses
    y = chart_y + chart_h + Inches(1.3)
    col2 = (CONTENT_W - GAP) / 2
    left, right = MARGIN, MARGIN + col2 + GAP
    yy = section_title(slide, left, y, 6, "Do the top papers stay on top?", col2)
    _stay_chart(slide, left, yy, col2, Inches(5.4))
    text(slide, left, yy + Inches(5.5), col2, Inches(1.1), [
        ([run("Share of runs in which a top-10 % paper is still in the top 10 %, at 20 % damage. "
              "Under added noise most of PageRank’s top papers are fragile; the forest keeps more of them.",
              size=24, color=MUTED)], {}),
    ], name="stay caption")

    yy = section_title(slide, right, y, 7, "The hypotheses", col2)
    _hypotheses_table(slide, right, yy, col2)

    # ---------- row 5: lessons + conclusion
    y = yy + Inches(7.1)
    yy = section_title(slide, left, y, 8, "What the data taught us", col2)
    lessons = [
        ("Ties look like stability. ", "Betweenness is 0 for 91 % of papers, so over all papers it looks the "
                                       "most stable — yet its top 10 % is the least stable. Compare at the top."),
        ("Trees cannot extrapolate. ", "After damage, up to 60 % of the inputs leave the training range and get "
                                       "one value. Part of the forest’s stability is blindness."),
        ("One citation, 1,545 places. ", "Paper 15004560 fell from rank 67 to 1,612 after gaining one reference: "
                                         "its betweenness jumped from 0."),
    ]
    text(slide, left, yy, col2, Inches(3.8), [
        ([run(b, size=25, bold=True, color=INK), run(t, size=25)], {"bullet": True, "after": 10, "line": 1.0})
        for b, t in lessons
    ], name="lessons")

    box(slide, right, y, col2, Inches(5.25), INK, MSO_SHAPE.ROUNDED_RECTANGLE, name="conclusion panel")
    text(slide, right + Inches(0.6), y + Inches(0.45), col2 - Inches(1.2), Inches(4.4), [
        ([run("Conclusion", font=HEAD_FONT, size=46, bold=True, color=WHITE)], {"after": 14}),
        ([run("A learned combination of centralities is more robust than PageRank against noisy citations, "
              "especially at the top of the ranking. The gain does not extend to missing citations, and simple "
              "local measures stay the most stable.", size=30, color=WHITE)], {"after": 18, "line": 1.05}),
        ([run("Limits: ", size=26, bold=True, color="F2C57C"),
          run("one dataset, one target (PageRank), one model type; trees cannot extrapolate.", size=26,
              color="CADCE8")], {}),
    ], name="conclusion")

    # ---------- footer
    text(slide, MARGIN, H - Inches(1.05), CONTENT_W, Inches(0.6), [
        ([run("Data: PubMed-Diabetes (Namata et al., 2012; LINQS).  Tools: Python, NetworkX, igraph, scikit-learn.  "
              f"Code, notebooks and interactive demo: {REPO_URL.replace('https://', '')}", size=20, color=MUTED)],
         {}),
    ], name="footer")

    out = ROOT / "poster" / "poster.pptx"
    prs.save(out)
    print("saved", out)


def _legend(slide, x, y, w, h):
    items = list(METHOD_STYLE.values())
    lx = x
    for i, (label, color, width) in enumerate(items):
        # line + label (about 0.13 in per character at 23 pt) + a fixed gap to the next entry
        item_w = Inches(1.0) + Inches(0.13) * len(label) + Inches(1.2)
        ln = slide.shapes.add_connector(1, *emu(lx, y + h / 2, lx + Inches(0.8), y + h / 2))
        ln.line.color.rgb = rgb(color)
        ln.line.width = Pt(width + 1)
        ln.name = f"legend line {i + 1}"
        text(slide, lx + Inches(1.0), y, item_w - Inches(1.1), h, [
            ([run(label, size=23, bold=i == 0, color=TEXT)], {})], anchor=MSO_ANCHOR.MIDDLE,
            name=f"legend label {i + 1}")
        lx += item_w


def _spearman_chart(slide, x, y, w, h, kind, title):
    data = XyChartData()
    for method in DRAW_ORDER:
        s = data.add_series(METHOD_STYLE[method][0])
        for level in [0.05, 0.10, 0.20]:
            s.add_data_point(level * 100, spearman(kind, method, level))
    frame = slide.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_LINES, *emu(x, y, w, h), data)
    frame.name = f"chart spearman {kind}"
    chart = frame.chart
    chart.has_legend = False
    chart.font.name, chart.font.size = BODY_FONT, Pt(22)
    chart.has_title = True
    chart.chart_title.text_frame.text = title
    tp = chart.chart_title.text_frame.paragraphs[0]
    tp.runs[0].font.size, tp.runs[0].font.bold = Pt(26), True
    tp.runs[0].font.color.rgb = rgb(INK)

    va, ca = chart.value_axis, chart.category_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 0.3, 1.0, 0.1
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = "0.0", False
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(LINE)
    ca.minimum_scale, ca.maximum_scale, ca.major_unit = 0, 25, 5
    ca.tick_labels.number_format, ca.tick_labels.number_format_is_linked = '0"%"', False
    ca.has_major_gridlines = False
    style_axis_text(va)
    style_axis_text(ca)

    for series, method in zip(chart.series, DRAW_ORDER):
        _, color, width = METHOD_STYLE[method]
        series.smooth = False
        series.format.line.color.rgb = rgb(color)
        series.format.line.width = Pt(width)
        series.marker.style = XL_MARKER_STYLE.CIRCLE
        series.marker.size = 14 if method == "rf_without_pagerank" else 10
        series.marker.format.fill.solid()
        series.marker.format.fill.fore_color.rgb = rgb(color)
        series.marker.format.line.color.rgb = rgb(color)


def _stay_chart(slide, x, y, w, h):
    data = CategoryChartData()
    data.categories = ["Removed", "Added, newer → older", "Added at random"]
    for method in METHOD_STYLE:
        data.add_series(METHOD_STYLE[method][0],
                        [stay_top(k, method) for k in ["remove", "add_dag", "add_random"]])
    frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, *emu(x, y, w, h), data)
    frame.name = "chart stay in top"
    chart = frame.chart
    chart.has_legend = False
    chart.font.name, chart.font.size = BODY_FONT, Pt(22)
    plot = chart.plots[0]
    plot.gap_width, plot.overlap = 55, -5
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.number_format, labels.number_format_is_linked = "0%", False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    labels.font.size = Pt(17)
    labels.font.color.rgb = rgb(TEXT)
    va = chart.value_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 0, 1.0, 0.25
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = "0%", False
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(LINE)
    style_axis_text(va)
    style_axis_text(chart.category_axis, size=23)
    for series, method in zip(chart.series, METHOD_STYLE):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = rgb(METHOD_STYLE[method][1])
        series.format.line.fill.background()


def _hypotheses_table(slide, x, y, w):
    rows = [
        ("Hypothesis", "Verdict", "What we found"),
        ("H1  Learned rankings are more stable than single measures",
         "Partly", "Better than PageRank under added noise; worse under removal; never better than degree or k-core"),
        ("H2  The forest absorbs a shock to one feature",
         "Not supported", "The forest with PageRank input relies on it and collapses with it — less stable "
                          "than PageRank itself"),
        ("H3  Key papers keep a high rank",
         "Supported", "Core papers stay in the top 10 % in 65–94 % of runs (adjusted core rule)"),
    ]
    verdict_color = {"Partly": "D98E04", "Not supported": "B23A48", "Supported": "3E7D5A"}
    shape = slide.shapes.add_table(len(rows), 3, *emu(x, y, w, Inches(6.4)))
    shape.name = "hypotheses table"
    table = shape.table
    widths = [0.36, 0.2, 0.44]
    for i, share in enumerate(widths):
        table.columns[i].width = int(w * share)
    table.rows[0].height = Inches(0.8)
    for r in range(1, len(rows)):
        table.rows[r].height = Inches(1.85)
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(INK if r == 0 else (PANEL if r % 2 else WHITE))
            cell.margin_left = cell.margin_right = Inches(0.2)
            cell.margin_top = cell.margin_bottom = Inches(0.12)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            rn = p.add_run()
            rn.text = keep_together(value)
            rn.font.name = BODY_FONT
            if r == 0:
                rn.font.size, rn.font.bold, rn.font.color.rgb = Pt(24), True, rgb(WHITE)
            elif c == 1:
                rn.font.size, rn.font.bold, rn.font.color.rgb = Pt(24), True, rgb(verdict_color[value])
            else:
                rn.font.size, rn.font.bold = Pt(22), c == 0
                rn.font.color.rgb = rgb(INK if c == 0 else TEXT)


if __name__ == "__main__":
    build()
